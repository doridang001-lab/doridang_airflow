import ast
import os
import time
import zipfile
from pathlib import Path

import openpyxl

from modules.extract import crawling_toorder_sales_report as target


def _write_minimal_xlsx(path: Path) -> None:
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types></Types>")
        zf.writestr("xl/workbook.xml", "<workbook></workbook>")


def test_cleanup_expected_report_downloads_removes_only_matching_report_files(tmp_path):
    expected = Path(target.ORIG_DATEDETAIL_FILENAME).stem
    old_xlsx = tmp_path / f"{expected}.xlsx"
    old_partial = tmp_path / f"{expected}.xlsx.crdownload"
    unrelated = tmp_path / "다른보고서.xlsx"
    old_xlsx.write_text("old", encoding="utf-8")
    old_partial.write_text("partial", encoding="utf-8")
    unrelated.write_text("keep", encoding="utf-8")

    target._cleanup_expected_report_downloads(tmp_path, expected)

    assert not old_xlsx.exists()
    assert not old_partial.exists()
    assert unrelated.exists()


def test_wait_for_report_xlsx_download_ignores_stale_files(tmp_path):
    expected = Path(target.ORIG_DATEDETAIL_FILENAME).stem
    stale = tmp_path / f"{expected}_old.xlsx"
    fresh = tmp_path / f"{expected}_fresh.xlsx"
    _write_minimal_xlsx(stale)
    _write_minimal_xlsx(fresh)

    started_at = time.time()
    os.utime(stale, (started_at - 30, started_at - 30))
    os.utime(fresh, (started_at + 1, started_at + 1))

    assert (
        target._wait_for_report_xlsx_download(
            download_dir=tmp_path,
            expected_stem=expected,
            download_started_at=started_at,
            timeout_sec=8,
        )
        == fresh
    )


def test_download_dir_diagnostics_reports_recent_partial_download(tmp_path):
    started_at = time.time()
    partial = tmp_path / "downloads.html.crdownload"
    old_partial = tmp_path / "downloads.html (old).crdownload"
    partial.write_bytes(b"Cr24" + b"\0" * 12)
    old_partial.write_bytes(b"old")
    os.utime(partial, (started_at + 1, started_at + 1))
    os.utime(old_partial, (started_at - 60, started_at - 60))

    diagnostics = target._download_dir_diagnostics(tmp_path, started_at)

    assert any(
        "downloads.html.crdownload" in item and "43 72 32 34" in item
        for item in diagnostics
    )
    assert all("downloads.html (old).crdownload" not in item for item in diagnostics)


def test_report_browser_options_disable_default_chrome_extensions(tmp_path):
    options = target._build_report_browser_options("account", tmp_path)

    assert "--disable-extensions" in options.arguments
    assert "--disable-component-extensions-with-background-pages" in options.arguments
    assert "--disable-default-apps" in options.arguments
    assert "--disable-sync" in options.arguments
    assert options.experimental_options["prefs"]["download.default_directory"] == str(
        tmp_path.absolute()
    )
    assert options.experimental_options["prefs"]["download.extensions_to_open"] == ""


def test_daily_date_page_initializes_download_guard_in_function_scope():
    source = Path(target.__file__).read_text(encoding="utf-8")
    tree = ast.parse(source)
    func = next(
        node
        for node in tree.body
        if isinstance(node, ast.FunctionDef)
        and node.name == "run_crawling_daily_date_page"
    )

    assigned_names = {
        child.id
        for child in ast.walk(func)
        if isinstance(child, ast.Name) and isinstance(child.ctx, ast.Store)
    }
    called_names = {
        child.func.id
        for child in ast.walk(func)
        if isinstance(child, ast.Call) and isinstance(child.func, ast.Name)
    }

    assert "expected_stem" in assigned_names
    assert "_cleanup_expected_report_downloads" in called_names
    assert "_download_datedetail_month" in called_names


def test_daily_date_page_reuses_datedetail_downloader(tmp_path, monkeypatch):
    class Driver:
        current_url = "https://ceo.toorder.co.kr/dashboard/sales-report/datedetail"

        def get(self, _url):
            self.current_url = "https://ceo.toorder.co.kr/dashboard/sales-report/datedetail"

        def quit(self):
            pass

    monkeypatch.setattr(target, "PIPELINE_RETRIES", 1)
    monkeypatch.setattr(target, "_cleanup_expected_report_downloads", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(target, "_launch_report_browser", lambda *_args, **_kwargs: Driver())
    monkeypatch.setattr(target, "_do_login", lambda *_args, **_kwargs: True)
    monkeypatch.setattr(target, "_open_datedetail_report_tab", lambda *_args, **_kwargs: True)
    monkeypatch.setattr(target, "_set_date_range", lambda *_args, **_kwargs: True)

    def _fake_download(_driver, _account, month_start, month_end, _download_dir):
        report_path = tmp_path / "report.xlsx"
        _write_minimal_xlsx(report_path)
        return {
            "success": True,
            "file": str(report_path),
            "month": month_start[:7],
            "error": None,
            "range": f"{month_start}~{month_end}",
        }

    monkeypatch.setattr(target, "_download_datedetail_month", _fake_download)

    result = target.run_crawling_daily_date_page(
        toorder_id="account",
        toorder_pw="pw",
        target_date="2026-08-09",
        download_dir=tmp_path,
    )

    assert result["success"] is True
    assert result["date"] == "2026-08-09"
    assert result["range"] == "2026-08-09~2026-08-09"
    assert Path(result["file"]).name.endswith("_260809.xlsx")


def test_daily_date_page_retries_datedetail_download_failure(tmp_path, monkeypatch):
    class Driver:
        current_url = "https://ceo.toorder.co.kr/dashboard/sales-report/datedetail"

        def get(self, _url):
            self.current_url = "https://ceo.toorder.co.kr/dashboard/sales-report/datedetail"

        def quit(self):
            pass

    calls = []

    monkeypatch.setattr(target, "PIPELINE_RETRIES", 2)
    monkeypatch.setattr(target, "PIPELINE_RETRY_BASE_SEC", 0)
    monkeypatch.setattr(target, "_cleanup_expected_report_downloads", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(target, "_launch_report_browser", lambda *_args, **_kwargs: Driver())
    monkeypatch.setattr(target, "_do_login", lambda *_args, **_kwargs: True)
    monkeypatch.setattr(target, "_open_datedetail_report_tab", lambda *_args, **_kwargs: True)
    monkeypatch.setattr(target, "_set_date_range", lambda *_args, **_kwargs: True)
    monkeypatch.setattr(target.time, "sleep", lambda *_args, **_kwargs: None)

    def _fake_download(_driver, _account, month_start, month_end, _download_dir):
        calls.append((month_start, month_end))
        if len(calls) == 1:
            return {
                "success": False,
                "file": None,
                "month": month_start[:7],
                "error": "downloaded file not found",
            }
        report_path = tmp_path / "report.xlsx"
        _write_minimal_xlsx(report_path)
        return {
            "success": True,
            "file": str(report_path),
            "month": month_start[:7],
            "error": None,
        }

    monkeypatch.setattr(target, "_download_datedetail_month", _fake_download)

    result = target.run_crawling_daily_date_page(
        toorder_id="account",
        toorder_pw="pw",
        target_date="2026-08-07",
        download_dir=tmp_path,
    )

    assert result["success"] is True
    assert result["date"] == "2026-08-07"
    assert calls == [("2026-08-07", "2026-08-07"), ("2026-08-07", "2026-08-07")]


def test_daily_date_page_retries_retriable_login_failure(tmp_path, monkeypatch):
    class Driver:
        current_url = "https://ceo.toorder.co.kr/dashboard/sales-report/datedetail"

        def get(self, _url):
            self.current_url = "https://ceo.toorder.co.kr/dashboard/sales-report/datedetail"

        def quit(self):
            pass

    login_calls = []

    monkeypatch.setattr(target, "PIPELINE_RETRIES", 2)
    monkeypatch.setattr(target, "PIPELINE_RETRY_BASE_SEC", 0)
    monkeypatch.setattr(target, "_cleanup_expected_report_downloads", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(target, "_launch_report_browser", lambda *_args, **_kwargs: Driver())
    monkeypatch.setattr(target, "_open_datedetail_report_tab", lambda *_args, **_kwargs: True)
    monkeypatch.setattr(target, "_set_date_range", lambda *_args, **_kwargs: True)
    monkeypatch.setattr(target.time, "sleep", lambda *_args, **_kwargs: None)

    def _fake_login(driver, *_args, **_kwargs):
        login_calls.append(1)
        if len(login_calls) == 1:
            setattr(driver, "_toorder_login_error", "React 앱 로드 타임아웃")
            return False
        return True

    def _fake_download(_driver, _account, month_start, _month_end, _download_dir):
        report_path = tmp_path / "report.xlsx"
        _write_minimal_xlsx(report_path)
        return {
            "success": True,
            "file": str(report_path),
            "month": month_start[:7],
            "error": None,
        }

    monkeypatch.setattr(target, "_do_login", _fake_login)
    monkeypatch.setattr(target, "_download_datedetail_month", _fake_download)

    result = target.run_crawling_daily_date_page(
        toorder_id="account",
        toorder_pw="pw",
        target_date="2026-08-08",
        download_dir=tmp_path,
    )

    assert result["success"] is True
    assert len(login_calls) == 2


def test_daily_date_page_recovers_late_existing_datedetail_file(tmp_path, monkeypatch):
    expected_stem = Path(target.ORIG_DATEDETAIL_FILENAME).stem
    late_file = tmp_path / f"{expected_stem}.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "종합"
    ws.cell(1, 1).value = "종합보고서_일별상세_매출보고서"
    wb.create_sheet("2026-08-07")
    wb.save(late_file)

    monkeypatch.setattr(target, "_launch_report_browser", lambda *_args, **_kwargs: None)

    result = target.run_crawling_daily_date_page(
        toorder_id="account",
        toorder_pw="pw",
        target_date="2026-08-07",
        download_dir=tmp_path,
    )

    assert result["success"] is True
    assert Path(result["file"]).name == f"{expected_stem}_260807.xlsx"
    assert not late_file.exists()
