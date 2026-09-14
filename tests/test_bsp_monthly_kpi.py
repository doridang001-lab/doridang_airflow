from __future__ import annotations

from datetime import date

import pandas as pd
from openpyxl import Workbook, load_workbook

from modules.transform.pipelines.db import DB_Bsp_Monthly_Kpi as pipeline


def _write_monthly_kpi(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    headers = [
        "주시작일",
        "ym",
        "주차",
        "담당자",
        "송파삼전점 영수건수",
        "송파삼전점 영수건수 목표",
        "가맹점 주문건수",
        "가맹점 주문건수 목표",
        "인스타그램 팔로워",
        "인스타그램 팔로워 목표",
        "카카오채널 친구",
        "카카오채널 친구 목표",
    ]
    ws.append(headers)
    ws.append([date(2026, 8, 3), "2026_08", "8월 2주차", "오나영", None, 1600, None, 62000, None, 2000, None, 2000])
    ws.append([date(2026, 8, 10), "2026_08", "8월 3주차", "오나영", None, 1600, None, 62000, None, 2000, None, 2000])
    ws.append([date(2026, 8, 17), "2026_08", "8월 4주차", "오나영", 999, 1600, 999, 62000, 999, 2000, 999, 2000])
    wb.save(path)


def test_snapshot_latest_is_keyed_by_week_start(tmp_path, monkeypatch):
    csv_path = tmp_path / "instagram_snapshot.csv"
    pd.DataFrame(
        [
            {"collect_date": "2026-08-07", "account": "doridang_official", "followers": 1409},
            {"collect_date": "2026-08-10", "account": "doridang_official", "followers": 1411},
            {"collect_date": "2026-08-12", "account": "doridang_official", "followers": 1412},
        ]
    ).to_csv(csv_path, index=False, encoding="utf-8-sig")

    result = pipeline._load_snapshot_latest_by_week_start(
        csv_path,
        id_col="account",
        id_value="doridang_official",
        value_col="followers",
    )

    assert result[date(2026, 8, 3)] == 1409
    assert result[date(2026, 8, 10)] == 1412


def test_sync_monthly_kpi_writes_week_specific_values_and_clears_missing(tmp_path, monkeypatch):
    xlsx_path = tmp_path / "monthly_kpi.xlsx"
    instagram_path = tmp_path / "instagram_snapshot.csv"
    kakao_path = tmp_path / "kakao_friends.csv"
    _write_monthly_kpi(xlsx_path)

    pd.DataFrame(
        [
            {"collect_date": "2026-08-07", "account": "doridang_official", "followers": 1409},
            {"collect_date": "2026-08-10", "account": "doridang_official", "followers": 1411},
        ]
    ).to_csv(instagram_path, index=False, encoding="utf-8-sig")
    pd.DataFrame(
        [
            {"collect_date": "2026-08-07", "channel_id": "_UxiaxiG", "friends": 1044},
            {"collect_date": "2026-08-10", "channel_id": "_UxiaxiG", "friends": 1045},
        ]
    ).to_csv(kakao_path, index=False, encoding="utf-8-sig")

    sales = pd.DataFrame(
        [
            {"week_start": date(2026, 8, 3), "hall_order_cnt": 163, "total_order_cnt": 15907},
            {"week_start": date(2026, 8, 10), "hall_order_cnt": 54, "total_order_cnt": 6000},
        ]
    )
    monkeypatch.setattr(pipeline, "BSP_MONTHLY_KPI_XLSX", xlsx_path)
    monkeypatch.setattr(pipeline, "INSTAGRAM_SNAPSHOT_CSV_PATH", instagram_path)
    monkeypatch.setattr(pipeline, "KAKAO_FRIENDS_CSV_PATH", kakao_path)
    monkeypatch.setattr(pipeline, "_load_sales_by_week_start", lambda: sales)

    result = pipeline.sync_monthly_kpi()

    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb["Sheet1"]
    assert "OK: monthly_kpi.xlsx 변경" in result
    assert [ws.cell(2, col).value for col in (5, 7, 9, 11)] == [163, 15907, 1409, 1044]
    assert [ws.cell(3, col).value for col in (5, 7, 9, 11)] == [54, 6000, 1411, 1045]
    assert [ws.cell(4, col).value for col in (5, 7, 9, 11)] == [None, None, None, None]
