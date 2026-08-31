"""
OKPOS product export download/save pipeline.
"""

import logging
import os
import shutil
import time
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile, is_zipfile

from openpyxl import load_workbook
from selenium.common.exceptions import TimeoutException
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

from modules.transform.pipelines.db.DB_OKPOS_Sales import (  # noqa: E402
    DOWNLOAD_TIMEOUT,
    WAIT_TIMEOUT,
    _launch_browser,
    _login,
    _setup_download_dir,
    _wait_for_download,
)
from modules.transform.utility.paths import ANALYTICS_DB, TEMP_DIR

logger = logging.getLogger(__name__)

OKPOS_PRODUCT_URL = "https://my.okpos.co.kr/asp/base/prod/search"
OKPOS_PRODUCT_FILE_NAME = "상품조회.xlsx"
OKPOS_PRODUCT_PER_PAGE = "5000"
OKPOS_PRODUCT_CODE_HEADER = "상품코드"
OKPOS_PRODUCT_SEARCH_TIMEOUT = int(os.getenv("OKPOS_PRODUCT_SEARCH_TIMEOUT", "90"))


def _cleanup_product_downloads(download_dir: Path) -> None:
    download_dir.mkdir(parents=True, exist_ok=True)
    stem = Path(OKPOS_PRODUCT_FILE_NAME).stem
    allowed_suffixes = {".xlsx", ".crdownload", ".part", ".tmp"}
    for path in download_dir.glob(f"{stem}*"):
        try:
            if path.is_file() and path.suffix.lower() in allowed_suffixes:
                path.unlink(missing_ok=True)
        except Exception:
            continue


def _okpos_product_filename_matches(path: Path) -> bool:
    return Path(OKPOS_PRODUCT_FILE_NAME).stem in path.stem


def _validate_okpos_product_xlsx(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(f"OKPOS product xlsx not found: {path}")
    if not is_zipfile(path):
        raise ValueError(f"OKPOS product file is not a real Excel workbook or is corrupted: {path}")

    try:
        with ZipFile(path) as workbook_zip:
            if "[Content_Types].xml" not in workbook_zip.namelist():
                raise ValueError(f"OKPOS product file is not a real Excel workbook or is corrupted: {path}")
    except BadZipFile as exc:
        raise ValueError(f"OKPOS product file is not a real Excel workbook or is corrupted: {path}") from exc

    try:
        workbook = load_workbook(BytesIO(path.read_bytes()), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"OKPOS product file is not a real Excel workbook or is corrupted: {path}") from exc

    try:
        if not workbook.worksheets:
            raise ValueError(f"OKPOS product file has no worksheets: {path}")

        worksheet = workbook.worksheets[0]
        rows = worksheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration as exc:
            raise ValueError(f"OKPOS product file is empty: {path}") from exc

        header_values = [str(value).strip() if value is not None else "" for value in header]
        if OKPOS_PRODUCT_CODE_HEADER not in header_values:
            raise ValueError(
                f"OKPOS product file has no '{OKPOS_PRODUCT_CODE_HEADER}' header: "
                f"{path} | columns={header_values}"
            )

        code_idx = header_values.index(OKPOS_PRODUCT_CODE_HEADER)
        data_rows = 0
        product_code_rows = 0
        for row in rows:
            row_values = list(row)
            if any(value is not None and str(value).strip() != "" for value in row_values):
                data_rows += 1
            code_value = row_values[code_idx] if code_idx < len(row_values) else None
            if code_value is not None and str(code_value).strip() != "":
                product_code_rows += 1

        if product_code_rows == 0:
            raise ValueError(
                "OKPOS product file has no product rows: "
                f"{path} | data_rows={data_rows} | product_code_rows={product_code_rows}"
            )
    finally:
        workbook.close()


def _product_grid_state(driver) -> dict[str, int | str | None]:
    return driver.execute_script(
        """
        const codeCells = Array.from(document.querySelectorAll('td.HideCol0prodCd'));
        const productRows = codeCells.filter((cell) => (cell.innerText || '').trim() !== '').length;
        const loading = Array.from(document.querySelectorAll('*')).filter((el) => {
          const style = getComputedStyle(el);
          const text = (el.innerText || '').trim();
          return (el.offsetWidth || el.offsetHeight || el.getClientRects().length)
            && text.includes('로딩')
            && style.display !== 'none'
            && style.visibility !== 'hidden';
        }).length;
        return {
          productRows,
          loading,
          jqueryActive: window.jQuery ? window.jQuery.active : null,
          bodyText: (document.body && document.body.innerText || '').slice(0, 500),
        };
        """
    )


def _wait_for_product_search_results(driver, timeout: int = OKPOS_PRODUCT_SEARCH_TIMEOUT) -> dict:
    deadline = time.monotonic() + timeout
    last_state: dict = {}
    while time.monotonic() < deadline:
        last_state = _product_grid_state(driver)
        product_rows = int(last_state.get("productRows") or 0)
        loading = int(last_state.get("loading") or 0)
        jquery_active = last_state.get("jqueryActive")
        if product_rows > 0 and loading == 0 and jquery_active in (0, None):
            logger.info("OKPOS product search loaded: product_rows=%d", product_rows)
            return last_state
        time.sleep(1)

    raise TimeoutException(
        "OKPOS product search did not load product rows before export: "
        f"timeout={timeout}s | state={last_state}"
    )


def download_okpos_product(**context) -> str:
    download_dir = TEMP_DIR / "okpos_product_download"
    _cleanup_product_downloads(download_dir)

    driver = _launch_browser(download_dir=download_dir)
    try:
        wait = WebDriverWait(driver, WAIT_TIMEOUT)
        _setup_download_dir(driver, download_dir)
        _login(driver, wait)

        logger.info("Open OKPOS product page: %s", OKPOS_PRODUCT_URL)
        driver.get(OKPOS_PRODUCT_URL)
        time.sleep(2)

        per_page_el = wait.until(EC.presence_of_element_located((By.ID, "perPage")))
        driver.execute_script(
            "arguments[0].value = arguments[1];"
            "arguments[0].dispatchEvent(new Event('input',  {bubbles:true}));"
            "arguments[0].dispatchEvent(new Event('change', {bubbles:true}));",
            per_page_el,
            OKPOS_PRODUCT_PER_PAGE,
        )

        search_btn = wait.until(EC.element_to_be_clickable((By.ID, "search_send")))
        driver.execute_script("arguments[0].click();", search_btn)
        _wait_for_product_search_results(driver)

        existing_files = {path for path in download_dir.iterdir() if path.is_file()}
        export_btn = wait.until(
            EC.element_to_be_clickable(
                (
                    By.XPATH,
                    "//button[contains(@onclick,'exportSheet') or contains(.,'엑셀다운')]",
                )
            )
        )
        driver.execute_script("arguments[0].click();", export_btn)

        downloaded = _wait_for_download(
            download_dir,
            existing_files,
            timeout=DOWNLOAD_TIMEOUT,
            expected_suffixes={".xlsx"},
            filename_predicate=_okpos_product_filename_matches,
        )
        if downloaded is None:
            expected = download_dir / OKPOS_PRODUCT_FILE_NAME
            if expected.exists() and _okpos_product_filename_matches(expected):
                downloaded = expected
            else:
                raise TimeoutException(
                    f"OKPOS product download timed out after {DOWNLOAD_TIMEOUT}s. "
                    f"download_dir={download_dir}"
                )

        _validate_okpos_product_xlsx(downloaded)
        context["ti"].xcom_push(key="downloaded_path", value=str(downloaded))
        logger.info("OKPOS product downloaded: downloaded_file=%s", downloaded.name)
        return f"downloaded: {downloaded}"
    finally:
        try:
            driver.quit()
        except Exception:
            pass


def save_okpos_product(**context) -> str:
    downloaded_path = context["ti"].xcom_pull(task_ids="download_okpos_product", key="downloaded_path")
    if not downloaded_path:
        raise ValueError("download_okpos_product XCom(downloaded_path) is empty.")

    src = Path(str(downloaded_path))
    if not src.exists():
        raise FileNotFoundError(f"Downloaded product file not found: {src}")

    _validate_okpos_product_xlsx(src)

    dest_dir = ANALYTICS_DB / "okpos_product"
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / OKPOS_PRODUCT_FILE_NAME
    tmp = dest_dir / f"{dest.name}.tmp"

    try:
        shutil.copy2(src, tmp)
        _validate_okpos_product_xlsx(tmp)
        os.replace(tmp, dest)
        try:
            src.unlink(missing_ok=True)
        except Exception:
            pass
    finally:
        try:
            tmp.unlink(missing_ok=True)
        except Exception:
            pass

    context["ti"].xcom_push(key="saved_path", value=str(dest))
    logger.info(
        "OKPOS product saved: downloaded_file=%s | saved_file=%s",
        src.name,
        dest.name,
    )
    return f"saved: {dest}"
